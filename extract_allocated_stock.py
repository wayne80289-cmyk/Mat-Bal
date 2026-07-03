# -*- coding: utf-8 -*-
"""Extract Stock rows with Material Code from History Balance template workbooks (sample tooling only)."""

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HB_DIR = Path(__file__).resolve().parent / "History Balance"
OUT_PATH = HB_DIR / "Stock-Allocated-Material-Code-All-Months.xlsx"

EXPORT_COLS = [
    "月份",
    "來源檔案",
    "Stock工作表",
    "Material_Code",
    "OWNER",
    "MAT_SPEC",
    "T",
    "W",
    "L",
    "MAKER_CODE",
    "CUST_CODE",
    "FG_CODE",
    "MAT_NO",
    "P_REMAIN_WT_kg",
    "REMAIN_WT_kg",
    "MAT_KIND",
    "MILL_NO",
]


def _text(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    return str(v).strip()


def _num(v) -> float:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def find_stock_sheet(names: list[str]) -> str | None:
    for n in names:
        if n.strip().lower().startswith("stock"):
            return n
    return None


def build_col_map(headers: list[str]) -> dict[str, int]:
    m: dict[str, int] = {}
    for i, h in enumerate(headers):
        key = _text(h).lower().replace("\n", " ")
        if not key:
            continue
        if key in ("material code", "material code"):
            m["material_code"] = i
        elif key == "owner":
            m["owner"] = i
        elif key == "mat spec":
            m["mat_spec"] = i
        elif key == "t":
            m["t"] = i
        elif key == "w":
            m["w"] = i
        elif key == "l":
            m["l"] = i
        elif key == "maker code":
            m["maker_code"] = i
        elif key == "cust code":
            m["cust_code"] = i
        elif key == "fg code":
            m["fg_code"] = i
        elif key == "mat no":
            m["mat_no"] = i
        elif key == "p remain wt":
            m["p_remain_wt"] = i
        elif key == "remain wt":
            m["remain_wt"] = i
        elif key in ("mat kind", "kind"):
            m["mat_kind"] = i
        elif key == "mill no":
            m["mill_no"] = i
    return m


def extract_stock_allocated(path: Path, month_label: str) -> list[dict]:
    xl = pd.ExcelFile(path)
    stock_name = find_stock_sheet(xl.sheet_names)
    if not stock_name:
        return []

    raw = pd.read_excel(path, sheet_name=stock_name, header=None)
    hdr_row = None
    for i in range(min(20, len(raw))):
        row_text = " ".join(_text(v).lower() for v in raw.iloc[i].tolist())
        if "material" in row_text and "code" in row_text and "owner" in row_text:
            hdr_row = i
            break
    if hdr_row is None:
        return []

    headers = [_text(v) for v in raw.iloc[hdr_row].tolist()]
    cmap = build_col_map(headers)
    if "material_code" not in cmap:
        return []

    rows = []
    for i in range(hdr_row + 1, len(raw)):
        code = _text(raw.iloc[i, cmap["material_code"]])
        if not code or code in ("0", "Choose", "False"):
            continue

        def g(key, default=None):
            if key not in cmap:
                return default
            return raw.iloc[i, cmap[key]]

        rows.append({
            "月份": month_label,
            "來源檔案": path.name,
            "Stock工作表": stock_name,
            "Material_Code": code,
            "OWNER": _text(g("owner")),
            "MAT_SPEC": _text(g("mat_spec")),
            "T": _num(g("t")),
            "W": _num(g("w")),
            "L": _text(g("l")),
            "MAKER_CODE": _text(g("maker_code")),
            "CUST_CODE": _text(g("cust_code")),
            "FG_CODE": _text(g("fg_code")),
            "MAT_NO": _text(g("mat_no")),
            "P_REMAIN_WT_kg": _num(g("p_remain_wt")),
            "REMAIN_WT_kg": _num(g("remain_wt")),
            "MAT_KIND": _text(g("mat_kind")),
            "MILL_NO": _text(g("mill_no")),
        })
    return rows


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.alignment = wrap
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 10
        for row in range(1, min(ws.max_row + 1, 500)):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                max_len = max(max_len, min(len(str(val)) + 2, 50))
        ws.column_dimensions[letter].width = max_len
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    all_rows: list[dict] = []
    file_log = []

    for month_dir in sorted(HB_DIR.iterdir()):
        if not month_dir.is_dir() or month_dir.name.startswith("."):
            continue
        month_label = month_dir.name.replace(".2026", "/2026")
        for path in sorted(month_dir.glob("*.xlsx")):
            if path.name.startswith("~$"):
                continue
            if "CARRIER" in path.name.upper():
                file_log.append({"月份": month_label, "檔案": path.name, "狀態": "跳過（無 Material Code 欄）", "筆數": 0})
                continue
            try:
                rows = extract_stock_allocated(path, month_label)
                all_rows.extend(rows)
                file_log.append({"月份": month_label, "檔案": path.name, "狀態": "已擷取", "筆數": len(rows)})
            except Exception as e:
                file_log.append({"月份": month_label, "檔案": path.name, "狀態": f"錯誤: {e}", "筆數": 0})

    detail = pd.DataFrame(all_rows, columns=EXPORT_COLS)
    log_df = pd.DataFrame(file_log)

  # Summary by month + material code
    if not detail.empty:
        summary = (
            detail.groupby(["月份", "Material_Code", "MAT_SPEC"], as_index=False)
            .agg(
                筆數=("Material_Code", "count"),
                總重量_kg=("P_REMAIN_WT_kg", "sum"),
                總重量_噸=("P_REMAIN_WT_kg", lambda s: round(s.sum() / 1000, 3)),
                OWNER=("OWNER", "first"),
                T=("T", "first"),
                W=("W", "first"),
            )
            .sort_values(["月份", "Material_Code"])
        )
    else:
        summary = pd.DataFrame()

    ts = datetime.now().strftime("%Y-%m-%d %H:%M")
    with pd.ExcelWriter(OUT_PATH, engine="openpyxl") as writer:
        meta = pd.DataFrame([
            {"項目": "產出時間", "內容": ts},
            {"項目": "總筆數（有 Material Code）", "內容": len(detail)},
            {"項目": "涵蓋月份", "內容": ", ".join(sorted(detail["月份"].unique())) if not detail.empty else "—"},
            {"項目": "唯一 Material Code", "內容": detail["Material_Code"].nunique() if not detail.empty else 0},
        ])
        meta.to_excel(writer, sheet_name="說明", index=False)
        detail.to_excel(writer, sheet_name="明細_全部月份", index=False)
        summary.to_excel(writer, sheet_name="彙總_按月按Code", index=False)
        log_df.to_excel(writer, sheet_name="擷取紀錄", index=False)

    wb = load_workbook(OUT_PATH)
    for name in ("明細_全部月份", "彙總_按月按Code", "擷取紀錄", "說明"):
        if name in wb.sheetnames:
            style_sheet(wb[name])
    wb.save(OUT_PATH)

    print(f"Saved: {OUT_PATH}")
    print(f"  Rows: {len(detail)} | Unique codes: {detail['Material_Code'].nunique() if not detail.empty else 0}")
    for r in file_log:
        print(f"  {r['月份']} {r['檔案']}: {r['筆數']} rows")


if __name__ == "__main__":
    main()
